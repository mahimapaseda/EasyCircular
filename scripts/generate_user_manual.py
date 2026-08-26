#!/usr/bin/env python3
"""Generate a colourful EasyCircular English user-manual Word document.

Output: docs/user-manual/EasyCircular-User-Manual.docx
Images: docs/user-manual/images/

Usage:
    python scripts/generate_user_manual.py
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "user-manual" / "images"
OUTPUT = ROOT / "docs" / "user-manual" / "EasyCircular-User-Manual.docx"

# Brand palette aligned with the EasyCircular UI
NAVY = "0F172A"
NAVY_MID = "1E293B"
CYAN = "0891B2"
CYAN_DK = "0E7490"
BLUE = "2563EB"
EMERALD = "059669"
AMBER = "D97706"
FUCHSIA = "C026D3"
ROSE = "E11D48"
PURPLE = "7C3AED"
WHITE = "FFFFFF"
SLATE = "334155"
MUTED = "64748B"
ICE = "F0F9FF"
MINT = "ECFDF5"
CREAM = "FFFBEB"
BLUSH = "FFF1F2"
LILAC = "F5F3FF"
ROW_ALT = "F8FAFC"
BORDER = "E2E8F0"

FONT = "Calibri"
FONT_DISPLAY = "Calibri"


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run(run, *, size=11, bold=False, italic=False, color=SLATE, font=FONT, all_caps=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    run.font.all_caps = all_caps
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(key), font)


def shade_cell(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_border(cell, **sides) -> None:
    """sides: top/left/bottom/right = (val, sz, color) e.g. ('single', 8, CYAN)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = sides.get(edge, ("nil", 0, WHITE))
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), spec[0])
        el.set(qn("w:sz"), str(spec[1]))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec[2])
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_table_borders(table, color=BORDER, sz=4, val="single") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(borders)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_table_width(table, width_cm: float) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    twips = int(width_cm * 567)
    tblW.set(qn("w:w"), str(twips))
    tblW.set(qn("w:type"), "dxa")


def clear_cell(cell) -> None:
    cell.text = ""
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def cell_para(cell, text, *, size=11, bold=False, italic=False, color=SLATE, align="left", space_after=4, space_before=0):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
    if cell.paragraphs[0].text == "" and cell.paragraphs[0] is p:
        pass
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run(run, **kwargs)
    return run


def full_width_table(doc, cols=1, width_cm=17.0):
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, width_cm)
    set_table_borders(table, val="nil", sz=0)
    return table


def page_break(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def add_footer(doc) -> None:
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "12")
    top.set(qn("w:space"), "8")
    top.set(qn("w:color"), CYAN)
    pBdr.append(top)
    pPr.append(pBdr)

    add_text(p, "EasyCircular  ", size=9, bold=True, color=NAVY)
    add_text(p, "User manual     August 2026", size=9, color=MUTED)
    add_text(p, "          ", size=9)
    add_text(p, "Page ", size=9, color=MUTED)

    def add_field(paragraph, code):
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = code
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(fld_end)
        set_run(run, size=9, bold=True, color=CYAN_DK)

    add_field(p, " PAGE ")
    add_text(p, " of ", size=9, color=MUTED)
    add_field(p, " NUMPAGES ")


def cover_page(doc) -> None:
    banner = full_width_table(doc)
    cell = banner.cell(0, 0)
    shade_cell(cell, NAVY)
    set_cell_margins(cell, top=280, bottom=200, left=280, right=280)
    set_cell_border(cell, bottom=("single", 48, CYAN))
    clear_cell(cell)

    cell_para(cell, "EASYCIRCULAR", size=12, bold=True, color=CYAN, align="center", space_after=2)
    cell_para(cell, "User Manual", size=36, bold=True, color=WHITE, align="center", space_before=8, space_after=4)
    cell_para(
        cell,
        "A short guide for people in schools who have to\nread Ministry of Education circulars",
        size=13,
        color="CBD5E1",
        align="center",
        space_before=8,
        space_after=14,
    )
    cell_para(cell, "English   ·   August 2026", size=11, bold=True, color=CYAN, align="center", space_after=4)
    cell_para(cell, "Principals, teachers, office staff", size=10, italic=True, color="94A3B8", align="center", space_after=8)

    # Colour stripe
    stripe = full_width_table(doc, cols=4)
    colours = [CYAN, BLUE, EMERALD, AMBER]
    labels = ["Read the PDF", "Check the text", "Get a summary", "Keep the original"]
    for i, (fill, label) in enumerate(zip(colours, labels)):
        c = stripe.cell(0, i)
        shade_cell(c, fill)
        set_cell_margins(c, top=80, bottom=80, left=60, right=60)
        set_cell_border(c)
        clear_cell(c)
        cell_para(c, label, size=10, bold=True, color=WHITE, align="center", space_after=0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)

    # Four feature tiles
    tiles = full_width_table(doc, cols=2)
    features = [
        (ICE, CYAN_DK, "Reading the pages", "It pulls text out of ordinary PDFs and from scans, including Sinhala and Tamil."),
        (LILAC, PURPLE, "A short brief", "What the circular is about, who it is for, and a list of things the school has to do."),
        (MINT, EMERALD, "You check first", "Nothing is summarised until you have looked at the extracted text. OCR gets names wrong."),
        (CREAM, AMBER, "Ministry list", "You can import a file from moe.gov.lk instead of hunting for a copy in email."),
    ]
    for i, (bg, accent, title, body) in enumerate(features):
        r, col = divmod(i, 2)
        if r and len(tiles.rows) <= r:
            tiles.add_row()
        c = tiles.cell(r, col)
        shade_cell(c, bg)
        set_cell_margins(c, top=140, bottom=140, left=160, right=160)
        set_cell_border(c, top=("nil", 0, WHITE), bottom=("nil", 0, WHITE), left=("single", 24, accent), right=("nil", 0, WHITE))
        clear_cell(c)
        cell_para(c, title, size=12, bold=True, color=accent, space_after=4)
        cell_para(c, body, size=10, color=SLATE, space_after=0)

    note = full_width_table(doc)
    c = note.cell(0, 0)
    shade_cell(c, BLUSH)
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)
    set_cell_border(c, left=("single", 24, ROSE))
    clear_cell(c)
    cell_para(c, "Please read this", size=10, bold=True, color=ROSE, space_after=2)
    cell_para(
        c,
        "The summary is a shortcut, not the circular. If a date or a circular number looks wrong, open the PDF. If you imported from the catalog, the ministry page is still the official one.",
        size=10,
        color=SLATE,
        space_after=0,
    )


def heading(doc, number: str, title: str, color: str = CYAN_DK) -> None:
    table = full_width_table(doc)
    cell = table.cell(0, 0)
    shade_cell(cell, NAVY)
    set_cell_margins(cell, top=90, bottom=90, left=160, right=160)
    set_cell_border(cell, left=("single", 32, color), bottom=("single", 12, color))
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, f"{number}   ", size=14, bold=True, color=color)
    add_text(p, title, size=16, bold=True, color=WHITE)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(8)


def subheading(doc, text: str, color: str = BLUE) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    add_text(p, text, size=13, bold=True, color=NAVY)


def body(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, size=11, color=SLATE)


def bullets(doc, items: list[str], accent=CYAN_DK) -> None:
    for item in items:
        table = full_width_table(doc)
        cell = table.cell(0, 0)
        set_cell_margins(cell, top=40, bottom=40, left=140, right=80)
        set_cell_border(cell, left=("single", 16, accent))
        clear_cell(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, item, size=11, color=SLATE)


def callout(doc, kind: str, title: str, text: str) -> None:
    styles = {
        "tip": (MINT, EMERALD, "TIP"),
        "important": (CREAM, AMBER, "IMPORTANT"),
        "warning": (BLUSH, ROSE, "WARNING"),
        "note": (ICE, CYAN_DK, "NOTE"),
    }
    bg, accent, label = styles[kind]
    table = full_width_table(doc)
    cell = table.cell(0, 0)
    shade_cell(cell, bg)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_border(cell, left=("single", 28, accent))
    clear_cell(cell)
    cell_para(cell, f"{label}  ·  {title}", size=10, bold=True, color=accent, space_after=3)
    cell_para(cell, text, size=10, color=SLATE, space_after=0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def colored_table(doc, headers: list[str], rows: list[list[str]], header_fill=CYAN_DK) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 17.0)
    set_table_borders(table, color=BORDER, sz=4)
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        shade_cell(cell, header_fill)
        set_cell_margins(cell, top=70, bottom=70, left=100, right=80)
        clear_cell(cell)
        cell_para(cell, h, size=10, bold=True, color=WHITE, space_after=0)
    for i, row in enumerate(rows):
        prevent_row_split(table.rows[i + 1])
        fill = WHITE if i % 2 == 0 else ROW_ALT
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            shade_cell(cell, fill)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=80)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            clear_cell(cell)
            # First column often a button/label — bold it
            cell_para(cell, val, size=10, bold=(j == 0), color=NAVY if j == 0 else SLATE, space_after=0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def step_row(doc, steps: list[tuple[str, str, str, str]]) -> None:
    """steps: (number, title, body, fill_hex)."""
    table = full_width_table(doc, cols=len(steps))
    for i, (num, title, body_text, fill) in enumerate(steps):
        cell = table.cell(0, i)
        shade_cell(cell, fill)
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        set_cell_border(cell)
        clear_cell(cell)
        cell_para(cell, num, size=9, bold=True, color=WHITE, align="center", space_after=2)
        cell_para(cell, title, size=11, bold=True, color=WHITE, align="center", space_after=4)
        cell_para(cell, body_text, size=9, color="F8FAFC", align="center", space_after=0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)


def add_figure(doc, filename: str, caption: str, width_cm: float = 15.4) -> None:
    path = IMG / filename
    table = full_width_table(doc)
    cap_cell = table.cell(0, 0)
    shade_cell(cap_cell, NAVY)
    set_cell_margins(cap_cell, top=50, bottom=50, left=140, right=140)
    set_cell_border(cap_cell, bottom=("single", 12, CYAN))
    clear_cell(cap_cell)
    cell_para(cap_cell, caption, size=9, bold=True, color=WHITE, align="center", space_after=0)

    img_row = table.add_row()
    img_cell = img_row.cells[0]
    shade_cell(img_cell, "F1F5F9")
    set_cell_margins(img_cell, top=80, bottom=80, left=80, right=80)
    set_cell_border(img_cell, top=("nil", 0, WHITE), bottom=("single", 8, CYAN), left=("single", 4, BORDER), right=("single", 4, BORDER))
    clear_cell(img_cell)
    p = img_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    if path.exists():
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
    else:
        add_text(p, f"[Image not found: {filename}]", size=10, italic=True, color=ROSE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)


def add_two_figures(doc, left: tuple[str, str], right: tuple[str, str], width_cm: float = 7.3) -> None:
    table = full_width_table(doc, cols=2)
    for i, (filename, caption) in enumerate((left, right)):
        cell = table.cell(0, i)
        shade_cell(cell, NAVY)
        set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
        clear_cell(cell)
        cell_para(cell, caption, size=8, bold=True, color=WHITE, align="center", space_after=0)
    img_row = table.add_row()
    for i, (filename, _caption) in enumerate((left, right)):
        cell = img_row.cells[i]
        shade_cell(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        clear_cell(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        path = IMG / filename
        if path.exists():
            run = p.add_run()
            run.add_picture(str(path), width=Cm(width_cm))
        else:
            add_text(p, f"[Missing: {filename}]", size=9, italic=True, color=ROSE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)


def toc_page(doc) -> None:
    heading(doc, "00", "Contents", CYAN)
    items = [
        ("01", "What this is", "Why you would use it"),
        ("02", "What you need", "Browser, login, PDF size"),
        ("03", "The usual path", "Upload, check, summarise"),
        ("04", "Accounts", "Sign up, Google, sign out"),
        ("05", "Home", "The first screen"),
        ("06", "Uploading a PDF", "The drop box and the 50 MB limit"),
        ("07", "Checking the text", "Editor, then Text is correct"),
        ("08", "The summary", "What you get, and when it fails"),
        ("09", "Coloured labels in the text", "Dates, names, circular numbers"),
        ("10", "Saving and translating", "Export, Copy, language tabs"),
        ("11", "Library", "Your files"),
        ("12", "Official catalog", "Import from moe.gov.lk"),
        ("13", "On a phone", "Same account, smaller screen"),
        ("14", "Those little status labels", "Uploaded through Failed"),
        ("15", "When something goes wrong", "The usual snags"),
    ]
    colours = [CYAN, BLUE, PURPLE, FUCHSIA, EMERALD, AMBER, ROSE, CYAN_DK]
    table = full_width_table(doc, cols=1)
    # rebuild as multi-row
    for i, (num, title, hint) in enumerate(items):
        if i:
            table.add_row()
        row = table.rows[i]
        prevent_row_split(row)
        cell = row.cells[0]
        shade_cell(cell, WHITE if i % 2 == 0 else ROW_ALT)
        set_cell_margins(cell, top=70, bottom=70, left=140, right=140)
        set_cell_border(cell, left=("single", 20, colours[i % len(colours)]), bottom=("single", 4, BORDER))
        clear_cell(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, f"{num}    ", size=12, bold=True, color=colours[i % len(colours)])
        add_text(p, title, size=12, bold=True, color=NAVY)
        add_text(p, f"     {hint}", size=10, italic=True, color=MUTED)


def build() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.6)

    core = doc.core_properties
    core.title = "EasyCircular User Manual"
    core.author = "EasyCircular"
    core.subject = "Guide for school staff using EasyCircular"
    core.comments = "User manual with screenshots"

    add_footer(doc)

    cover_page(doc)
    page_break(doc)
    toc_page(doc)

    # 01
    page_break(doc)
    heading(doc, "01", "What this is", CYAN)
    body(
        doc,
        "If you work in a school and you get MOE circulars as PDFs, this is for you. You either drop in a file you already have, or you pull one from the ministry site. The system reads the pages (including scans in English, Sinhala, or Tamil), shows you the text, and then writes a short brief.",
    )
    body(
        doc,
        "You still have to look at that text before the summary runs. That is on purpose. OCR misreads names. Scans pick up stamps. Better you catch it than the summary invent a deadline.",
    )
    body(
        doc,
        "After that you can copy the brief, download it as a PDF, or switch it to Sinhala or Tamil if the original was in English (and the other way round).",
    )
    add_figure(doc, "home.png", "Figure 1  ·  Home. Start Processing if you have a file, Official circulars for the ministry list.")

    # 02
    heading(doc, "02", "What you need", BLUE)
    body(doc, "Chrome, Edge, Firefox, or Safari is fine. You need an account — email and a password, or Google. The file has to be a real PDF, not bigger than 50 MB.")
    body(doc, "You can look around without signing in. You cannot upload or import until you do.")
    callout(
        doc,
        "note",
        "Keep the original",
        "EasyCircular does not replace the circular, and it does not send anything to the Ministry on your behalf. File the PDF with whatever summary you export.",
    )

    # 03
    heading(doc, "03", "The usual path", EMERALD)
    body(doc, "Along the top of a document you will see four steps: Upload, Extract, Review, Summarize. That is the whole job, in order.")
    step_row(
        doc,
        [
            ("1", "Upload", "Sign in. Drop a PDF, or import one from the catalog.", CYAN_DK),
            ("2", "Extract", "Wait. Scanned pages take longer.", BLUE),
            ("3", "Review", "Fix the number and the dates. Then Text is correct.", PURPLE),
            ("4", "Summarize", "Read the brief. Export it if you need a copy.", EMERALD),
        ],
    )
    body(doc, "From home, Start Processing or New Analysis gets you a file in. Then: Editor, Save edits, Text is correct, Generate summary. The two things worth checking by hand are the circular number and the dates. Staff act on those.")
    add_figure(doc, "home-signed-in.png", "Figure 2  ·  Signed in — the upload box is available")

    # 04
    heading(doc, "04", "Accounts", FUCHSIA)
    subheading(doc, "New account", FUCHSIA)
    body(doc, "Click Get Started or Sign up. You will be asked for your name, email, a password of at least eight characters, and your job role (Teacher, Principal, or Education administration). District is optional; type a few letters and pick from the list.")
    body(doc, "Create account, or Continue with Google if that is easier. After that you land back at upload.")
    add_figure(doc, "sign-up.png", "Figure 3  ·  Sign-up form")

    subheading(doc, "Coming back", BLUE)
    body(doc, "Use Sign in. Same email and password, or Google again. Remember me is there if you want it.")
    callout(
        doc,
        "note",
        "Forgot password?",
        "The link is on the form. It does not actually reset anything yet. If you are stuck, try Google, or get someone to make a new account.",
    )
    add_figure(doc, "sign-in.png", "Figure 4  ·  Sign-in")
    body(doc, "Sign out is in the top bar on the home page, and in the left sidebar once you are in a document.")

    # 05
    heading(doc, "05", "Home", CYAN)
    body(doc, "This is the first screen. The dark Start Processing button takes you to upload (it will send you to sign in if you have not). Official circulars opens the moe.gov.lk list inside EasyCircular.")
    body(doc, "Get Started is sign in / sign up. The moon (or sun) icon switches dark mode. Library is your own files. The four cards under the fold are just a description of the product. You do not click them to do the work.")

    # 06
    heading(doc, "06", "Uploading a PDF", BLUE)
    body(doc, "If you are not signed in, the box says Sign in to upload. Fair enough.")
    body(doc, "From home, scroll down to the drop area, or once you are in the workspace use New Analysis on the left. Drop the file, or click to browse. You should see Uploading your circular…, then the document screen opens.")
    colored_table(
        doc,
        ["Limit", "What happens"],
        [
            ["Has to be a PDF", "Renaming a Word file to .pdf will not work"],
            ["50 MB", "Anything larger is rejected"],
            ["One file at a time", "There is no upload-this-whole-folder button"],
        ],
        BLUE,
    )
    add_figure(doc, "upload.png", "Figure 5  ·  The upload box on the home page")

    # 07
    heading(doc, "07", "Checking the text", PURPLE)
    body(doc, "Extraction usually starts on its own. If it does not, use Extract text. Re-extract runs it again if the first pass looked empty or garbled. Sinhala and Tamil scans can sit there for a while. Leave the tab open.")
    body(doc, "You cannot generate a summary until you have said the text is all right. That is the Review step.")
    bullets(
        doc,
        [
            "Open Source text. Editor is where you type. Highlights shows coloured tags later, after a summary exists.",
            "Correct the obvious stuff: circular number, dates, school names, anything OCR mangled in Sinhala or Tamil.",
            "Save edits (that also counts as reviewed). Then Text is correct, then Generate summary.",
        ],
        PURPLE,
    )
    callout(
        doc,
        "important",
        "Working copy",
        "The text on screen is a working copy. The PDF is still the circular. Fix the number and the dates here before you trust the brief.",
    )
    add_figure(doc, "review-text.png", "Figure 6  ·  Review. Read the text, then confirm it before you summarise.")

    # 08
    heading(doc, "08", "The summary", EMERALD)
    body(doc, "Generate summary shows Building brief… for a bit. What you get is a title, the number, who issued it, who it is for, when it takes effect, a purpose, the main requirements, Action items, and sometimes Key references.")
    colored_table(
        doc,
        ["Button", "What it is for"],
        [
            ["Edit", "Change the wording, then Save changes or Cancel. Original language only."],
            ["Copy", "Puts it on the clipboard"],
            ["Export", "Downloads a PDF of the brief — not the original circular"],
            ["Regenerate", "Runs it again from the text you have now"],
            ["Language tabs", "English / සිංහල / தமிழ், depending on the source"],
        ],
        EMERALD,
    )
    body(doc, "There is a small disclaimer on every brief. Read it. The model can skip things or make things up. If the file came from the catalog, Open on moe.gov.lk is a link back to the ministry page. Use it.")
    add_figure(doc, "summary.png", "Figure 7  ·  A finished brief. Action items are on the right.")

    subheading(doc, "When the summary is not quite right", ROSE)
    colored_table(
        doc,
        ["You will see", "What happened", "Sensible next step"],
        [
            ["AI summary unavailable", "The model did not run. You get a plainer extract instead.", "Try Regenerate in a few minutes."],
            ["Review suggested", "It thought a date was invented, or two circulars got mixed.", "Open the PDF. Do not act on the brief yet."],
            ["Unreadable brief", "The text came out as a mix of scripts.", "Regenerate."],
        ],
        ROSE,
    )

    # 09
    heading(doc, "09", "Coloured labels in the text", AMBER)
    body(doc, "After a summary, expand Source text and switch to Highlights. Dates, organisations, people, and legal-looking phrases are marked. You might see something like “13 entities · 90% confidence” on the bar. Treat that as a hint, not a score you report to anyone.")
    colored_table(
        doc,
        ["Colour", "Tag", "Usually"],
        [
            ["Yellow / amber", "DATE", "Deadlines and issue dates"],
            ["Blue", "ORG", "Ministry, a school, a provincial office"],
            ["Pink", "PERSON", "A name on the letter"],
            ["Teal", "LAW", "Circular numbers, acts"],
            ["Grey", "OTHER", "Anything else it tagged"],
        ],
        AMBER,
    )
    add_figure(doc, "entities.png", "Figure 8  ·  Highlights on the source text")

    # 10
    heading(doc, "10", "Saving and translating", CYAN_DK)
    subheading(doc, "Export", CYAN)
    body(doc, "Export is on the brief toolbar. That PDF is the summary, for a staff meeting or an email. If you only need a paste, use Copy.")
    subheading(doc, "Translate", BLUE)
    body(doc, "The language tabs are on the left of the same toolbar. English circulars usually offer Sinhala as well. A Sinhala or Tamil original will offer English. Click the tab and wait for Translating brief…")
    callout(
        doc,
        "tip",
        "Translations",
        "Handy for sharing with staff. Not more official than the circular you uploaded.",
    )
    add_figure(doc, "export-translate.png", "Figure 9  ·  Language tabs, Edit, Copy, Export")

    # 11
    heading(doc, "11", "Library", BLUE)
    body(doc, "Library in the sidebar (or from home) is your pile of work. Uploads and catalog imports both land here. The switch at the top is Your library versus Official catalog. On your library you can filter All, Summarized, or In progress, and you can search by file name or a bit of the summary.")
    body(doc, "Click a card to open it. Each card shows how old it is, pages, words, and how many tags it found. Imports from the ministry are marked Official moe.gov.lk. Delete is on the card. It asks once. After that the file is gone.")
    body(doc, "If you are not signed in you can still see the layout, but you will be asked to sign in before anything is saved against you.")
    add_figure(doc, "library.png", "Figure 10  ·  Library")

    # 12
    heading(doc, "12", "Official catalog", EMERALD)
    body(doc, "This is the list from moe.gov.lk, shown inside EasyCircular. Importing a PDF does not mean EasyCircular is the ministry. The summary is still made here, from that file.")
    body(doc, "Open Official catalog (home has Official circulars as well). Search. On an entry, open PDFs, then Import. If you are not signed in it will say Sign in to import. Importing… takes a moment. After that it is the same as an upload: extract, review, summarise. Previous and Next move through the list.")
    add_figure(doc, "official-catalog.png", "Figure 11  ·  Catalog search and import")

    # 13
    heading(doc, "13", "On a phone", PURPLE)
    body(doc, "It works in the phone browser with the same login. Buttons stack. New Analysis is still how you upload. If you are going to sit and correct a long Sinhala scan, a laptop is less painful.")
    add_two_figures(
        doc,
        ("mobile-home.png", "Figure 12  ·  Home on a phone"),
        ("mobile-summary.png", "Figure 13  ·  A brief on a phone"),
    )

    # 14
    heading(doc, "14", "Those little status labels", AMBER)
    colored_table(
        doc,
        ["Label", "Colour", "Meaning"],
        [
            ["Uploaded", "Cyan", "File is in. Text not out yet."],
            ["Ready to review", "Purple", "Text is there. You still have to check it."],
            ["Processing", "Amber", "Summary is running."],
            ["Completed", "Green", "Brief is ready."],
            ["Failed", "Red", "It broke. Try again or Retry."],
        ],
        AMBER,
    )

    # 15
    heading(doc, "15", "When something goes wrong", ROSE)
    body(doc, "Check the circular number and the dates first. If those are wrong, fix them in the editor and regenerate. Prefer Official catalog when you can, so you are not working from a photocopy someone emailed last year. Export once you are happy, and file that PDF with the original circular.")
    colored_table(
        doc,
        ["Problem", "Try this"],
        [
            ["Sign in to upload", "You are not logged in."],
            ["File rejected", "It is not a real PDF, or it is over 50 MB."],
            ["Blank or messy text", "Wait, then Re-extract. Edit what you can."],
            ["Summary is nonsense", "Fix the source, Save edits, Regenerate."],
            ["AI summary unavailable", "The model is down or busy. Try later."],
            ["Forgot password?", "Does not work. Use Google, or a password you still know."],
            ["A long scan is slow", "That is normal. Do not close the tab."],
        ],
        ROSE,
    )
    body(doc, "EasyCircular will not email the Ministry for you, will not file your staff return, and will not take a whole folder of PDFs in one go. If a date matters, read it in the original circular. The home page has a System Status strip. If that looks unhealthy, the summary may fail until someone brings the services back up.")

    # Close
    close = full_width_table(doc)
    cell = close.cell(0, 0)
    shade_cell(cell, NAVY)
    set_cell_margins(cell, top=200, bottom=200, left=240, right=240)
    set_cell_border(cell, top=("single", 24, CYAN))
    clear_cell(cell)
    cell_para(cell, "When in doubt, open the PDF", size=16, bold=True, color=WHITE, align="center", space_after=8)
    cell_para(
        cell,
        "This guide is for EasyCircular as it stands in August 2026. The System Status strip on the home page tells you if the site is actually up.",
        size=11,
        color="CBD5E1",
        align="center",
        space_after=0,
    )

    return doc


PDF_OUTPUT = OUTPUT.with_suffix(".pdf")


def export_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Export via Microsoft Word so PAGE fields resolve in the PDF."""
    import subprocess

    docx = str(docx_path.resolve()).replace("\\", "/")
    pdf = str(pdf_path.resolve()).replace("\\", "/")
    script = (
        "$ErrorActionPreference = 'Stop'\n"
        "$word = New-Object -ComObject Word.Application\n"
        "$word.Visible = $false\n"
        "$word.DisplayAlerts = 0\n"
        "try {\n"
        f"  $doc = $word.Documents.Open({repr(docx)})\n"
        "  $null = $doc.Fields.Update()\n"
        f"  $tmp = {repr(pdf.rsplit('.', 1)[0] + '-export.pdf')}\n"
        f"  $pdf = {repr(pdf)}\n"
        "  if (Test-Path $tmp) { Remove-Item $tmp -Force }\n"
        "  $doc.ExportAsFixedFormat($tmp, 17)\n"
        "  $doc.Close($false)\n"
        "  if (Test-Path $pdf) { Remove-Item $pdf -Force }\n"
        "  Move-Item $tmp $pdf\n"
        "} finally {\n"
        "  $word.Quit()\n"
        "  [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null\n"
        "}\n"
    )
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
        check=True,
    )


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    missing = [
        name
        for name in (
            "home.png",
            "home-signed-in.png",
            "sign-up.png",
            "sign-in.png",
            "upload.png",
            "review-text.png",
            "summary.png",
            "entities.png",
            "export-translate.png",
            "library.png",
            "official-catalog.png",
            "mobile-home.png",
            "mobile-summary.png",
        )
        if not (IMG / name).exists()
    ]
    if missing:
        print("Missing images:", ", ".join(missing))
    doc = build()
    doc.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")
    try:
        export_pdf(OUTPUT, PDF_OUTPUT)
        print(f"Wrote {PDF_OUTPUT}")
    except Exception as exc:
        print(f"PDF export failed: {exc}")
    print(f"Images: {IMG}")


if __name__ == "__main__":
    main()
