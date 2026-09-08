#!/usr/bin/env python3
"""Generate EasyCircular IEEE A4 conference paper from YAML + official template."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
CONF = ROOT / "docs" / "conference"
TEMPLATE = CONF / "conference-template-a4.docx"
CONTENT = CONF / "content.yaml"
OUTPUT = CONF / "EasyCircular-Conference-Paper.docx"
OUTPUT_BLIND = CONF / "EasyCircular-Conference-Paper-Blind.docx"
OUTPUT_PDF = CONF / "EasyCircular-Conference-Paper.pdf"

OOXML_PATCHES = [
    (
        "http://purl.oclc.org/ooxml/officeDocument/relationships/",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/",
    ),
    (
        "http://purl.oclc.org/ooxml/wordprocessingml/main",
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    ),
    (
        "http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing",
        "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    ),
    (
        "http://purl.oclc.org/ooxml/officeDocument/math",
        "http://schemas.openxmlformats.org/officeDocument/2006/math",
    ),
    ("http://purl.oclc.org/ooxml/drawingml/main", "http://schemas.openxmlformats.org/drawingml/2006/main"),
]

def load_content() -> dict[str, Any]:
    with CONTENT.open(encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def patch_template_ooxml(src: Path, dst: Path) -> None:
    shutil.copy2(src, dst)
    with zipfile.ZipFile(dst, "r") as zin:
        data = {n: zin.read(n) for n in zin.namelist()}
    for key in list(data):
        if not key.endswith((".xml", ".rels")):
            continue
        text = data[key].decode("utf-8")
        for old, new in OOXML_PATCHES:
            text = text.replace(old, new)
        data[key] = text.encode("utf-8")
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, blob in data.items():
            zout.writestr(name, blob)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def clear_paragraph(paragraph) -> None:
    paragraph.text = ""


def set_run_font(run, font_name: str, size_pt: int, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    rpr.append(sz_cs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


SINHALA_START, SINHALA_END, ZWJ = "\u0d80", "\u0dff", "\u200d"


def _script_of(ch: str) -> str | None:
    if SINHALA_START <= ch <= SINHALA_END or ch == ZWJ:
        return "si"
    if ch.isalpha():
        return "latin"
    return None


def split_by_script(text: str) -> list[tuple[str, str]]:
    segments: list[tuple[str, str]] = []
    script: str | None = None
    buf: list[str] = []
    for ch in text:
        ch_script = _script_of(ch)
        if ch_script is None or ch_script == script or script is None:
            if script is None and ch_script is not None:
                script = ch_script
            buf.append(ch)
            continue
        segments.append((script, "".join(buf)))
        script, buf = ch_script, [ch]
    if buf:
        segments.append((script or "latin", "".join(buf)))
    return segments


def add_runs(paragraph, text: str, meta: dict, size_pt: int, bold: bool = False, italic: bool = False) -> None:
    latin = meta.get("font_latin", "Times New Roman")
    sinhala = meta.get("font", "Iskoola Pota")
    for script, segment in split_by_script(text):
        if not segment:
            continue
        run = paragraph.add_run(segment)
        set_run_font(run, sinhala if script == "si" else latin, size_pt, bold=bold)
        if italic:
            run.font.italic = True


def set_paragraph_text(paragraph, text: str, meta: dict, size_pt: int, bold: bool = False) -> None:
    paragraph.text = ""
    add_runs(paragraph, text, meta, size_pt, bold=bold)


def replace_title(doc: Document, title: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name == "paper title":
            set_paragraph_text(paragraph, title, {"font_latin": "Times New Roman"}, 24, bold=True)
            return
    raise RuntimeError("paper title style not found in template")


def set_author_block(paragraph, author: dict, meta: dict) -> None:
    paragraph.text = ""
    lines = [
        author["name"],
        author["department"],
        f"({author['organization']})",
        author["city_country"],
        author["email"],
    ]
    for idx, line in enumerate(lines):
        if idx:
            paragraph.add_run().add_break()
        add_runs(paragraph, line, meta, 10)


def replace_authors(doc: Document, authors: list[dict], meta: dict, blind: bool) -> None:
    author2_placed = False
    for paragraph in doc.paragraphs:
        if not (paragraph.style and paragraph.style.name == "Author"):
            continue
        if blind:
            clear_paragraph(paragraph)
            continue
        text = paragraph.text
        if "1st Given Name" in text:
            set_author_block(paragraph, authors[0], meta)
        elif ("4th Given Name" in text or "3rd Given Name" in text) and not author2_placed and len(authors) > 1:
            set_author_block(paragraph, authors[1], meta)
            author2_placed = True
        else:
            clear_paragraph(paragraph)


def replace_abstract_keywords(doc: Document, content: dict, meta: dict) -> None:
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        if style == "Abstract":
            paragraph.text = ""
            add_runs(paragraph, "Abstract: ", meta, meta.get("font_size", 10), bold=True, italic=True)
            abstract = content.get("abstract_en") or content.get("abstract_si") or ""
            add_runs(paragraph, abstract, meta, meta.get("font_size", 10))
        elif style == "Keywords":
            paragraph.text = ""
            add_runs(paragraph, "Keywords: ", meta, meta.get("font_size", 10), italic=True)
            add_runs(paragraph, content["keywords_en"], meta, meta.get("font_size", 10))


def truncate_after_keywords(doc: Document) -> None:
    """Remove all template body text after the Keywords paragraph."""
    seen_keywords = False
    for paragraph in list(doc.paragraphs):
        style = paragraph.style.name if paragraph.style else ""
        if style == "Keywords":
            seen_keywords = True
            continue
        if seen_keywords:
            delete_paragraph(paragraph)
    for table in list(doc.tables):
        table._element.getparent().remove(table._element)


def add_body_paragraph(doc: Document, text: str, meta: dict, style: str | None = "Body Text") -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_after = Pt(3)
    add_runs(paragraph, text, meta, meta.get("font_size", 10))


def add_section_heading(doc: Document, heading: str, meta: dict) -> None:
    is_sub = heading[:2] in {"A.", "B.", "C.", "D."}
    style = "Heading 2" if is_sub else "Heading 1"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(6 if not is_sub else 4)
    paragraph.paragraph_format.space_after = Pt(3)
    add_runs(paragraph, heading, meta, meta.get("font_size", 10), bold=True)


def add_figure(doc: Document, fig_num: int, fig: dict, meta: dict) -> None:
    image_path = ROOT / fig["file"]
    if not image_path.exists():
        print(f"  warning: missing figure {image_path}")
        return
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    width = fig.get("width_cm", 7.5)
    run.add_picture(str(image_path), width=Cm(width))
    caption = doc.add_paragraph(style="figure caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(caption, f"Fig. {fig_num}. ", meta, 8, bold=True)
    add_runs(caption, fig["caption_si"], meta, 8)


def add_table(doc: Document, table_def: dict, meta: dict) -> None:
    marker = table_def.get("marker", "TABLE I")
    caption = doc.add_paragraph(style="figure caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(caption, f"{marker}. ", meta, 8, bold=True)
    add_runs(caption, table_def["caption_si"], meta, 8)
    headers = table_def["headers"]
    rows = table_def["rows"]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        table.style = "Normal Table"
    except KeyError:
        pass
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = ""
        add_runs(cell.paragraphs[0], header, meta, 8, bold=True)
    for row_idx, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            cell = table.rows[row_idx].cells[col]
            cell.text = ""
            add_runs(cell.paragraphs[0], value, meta, 8)


def add_references(doc: Document, references: list[str], meta: dict) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    add_runs(heading, "REFERENCES", meta, meta.get("font_size", 10), bold=True)
    for idx, ref in enumerate(references, start=1):
        paragraph = doc.add_paragraph(style="Body Text")
        paragraph.paragraph_format.left_indent = Cm(0.35)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.paragraph_format.space_after = Pt(2)
        add_runs(paragraph, f"[{idx}] ", meta, 8)
        add_runs(paragraph, ref, meta, 8)


def _mentions_fig(text: str, number: int) -> bool:
    return re.search(rf"Fig\.\s*{number}(?!\d)", text) is not None


def _mentions_table(text: str, marker: str) -> bool:
    escaped = re.escape(marker)
    return re.search(rf"{escaped}(?![IVX])", text) is not None


def append_body(doc: Document, content: dict, meta: dict) -> None:
    figures = sorted(content.get("figures") or [], key=lambda f: int(f.get("number", 0)), reverse=True)
    tables = sorted(content.get("tables") or [], key=lambda t: len(t.get("marker", "")), reverse=True)
    used_figs: set[str] = set()
    used_tables: set[str] = set()

    for section in content.get("sections") or []:
        add_section_heading(doc, section["heading_en"], meta)
        for paragraph_text in section.get("paragraphs_si") or []:
            add_body_paragraph(doc, paragraph_text, meta)
            for fig in figures:
                fid = fig["id"]
                if fid in used_figs:
                    continue
                if _mentions_fig(paragraph_text, int(fig["number"])):
                    add_figure(doc, int(fig["number"]), fig, meta)
                    used_figs.add(fid)
            for table_def in tables:
                tid = table_def["id"]
                if tid in used_tables:
                    continue
                if _mentions_table(paragraph_text, table_def["marker"]):
                    add_table(doc, table_def, meta)
                    used_tables.add(tid)

    add_references(doc, content.get("references") or [], meta)


def export_pdf(docx_path: Path, pdf_path: Path) -> None:
    if pdf_path.exists():
        try:
            pdf_path.unlink()
        except OSError:
            pass
    docx = str(docx_path.resolve()).replace("\\", "/")
    pdf = str(pdf_path.resolve()).replace("\\", "/")
    ps = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
  $doc = $word.Documents.Open("{docx}")
  $doc.ExportAsFixedFormat("{pdf}", 17)
  $doc.Close($false)
}} finally {{
  $word.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}}
"""
    import subprocess as _sp

    result = _sp.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        print("  PDF export skipped (Microsoft Word required)")
        if result.stderr:
            print(f"  {result.stderr.strip()}")


def build_document(content: dict, blind: bool, output_path: Path) -> None:
    meta = content.get("meta") or {}
    meta.setdefault("font", "Iskoola Pota")
    meta.setdefault("font_latin", "Times New Roman")
    meta.setdefault("font_size", 10)

    work = CONF / "_working-template.docx"
    patch_template_ooxml(TEMPLATE, work)
    doc = Document(str(work))

    replace_title(doc, content["title_en"])
    replace_authors(doc, content.get("authors") or [], meta, blind=blind)
    replace_abstract_keywords(doc, content, meta)
    truncate_after_keywords(doc)
    append_body(doc, content, meta)

    saved = output_path
    try:
        doc.save(str(output_path))
    except PermissionError:
        saved = output_path.with_name(output_path.stem + "-updated.docx")
        doc.save(str(saved))
        print(f"  Note: {output_path.name} is open in Word; saved as {saved.name}")
    work.unlink(missing_ok=True)
    print(f"Wrote {saved}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate EasyCircular IEEE conference paper")
    parser.add_argument("--blind", action="store_true", help="Omit author names (double-blind review)")
    parser.add_argument("--no-pdf", action="store_true", help="Skip PDF export")
    args = parser.parse_args()

    if not TEMPLATE.exists():
        print(f"Missing template: {TEMPLATE}")
        return 1
    if not CONTENT.exists():
        print(f"Missing content: {CONTENT}")
        return 1

    content = load_content()
    subprocess.run(
        ["python", str(ROOT / "scripts" / "render_conference_survey_chart.py")],
        check=False,
    )
    output = OUTPUT_BLIND if args.blind else OUTPUT
    build_document(content, blind=args.blind, output_path=output)

    if not args.no_pdf and not args.blind:
        export_pdf(output, OUTPUT_PDF)
        if OUTPUT_PDF.exists() and OUTPUT_PDF.stat().st_size > 0:
            print(f"Wrote {OUTPUT_PDF}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
