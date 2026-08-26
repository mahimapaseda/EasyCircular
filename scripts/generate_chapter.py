#!/usr/bin/env python3
"""Generate a standalone thesis chapter Word document from YAML content."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import generate_thesis as gt  # noqa: E402

CHAPTER_MAP = {
    "1": ("ch01_introduction.yaml", "ch01_introduction", "Chapter-01-Introduction.docx"),
    "2": ("ch02_literature.yaml", "ch02_literature", "Chapter-02-Literature.docx"),
    "3": ("ch03_design.yaml", "ch03_design", "Chapter-03-Design.docx"),
    "4": ("ch04_development.yaml", "ch04_development", "Chapter-04-Development.docx"),
    "5": ("ch05_evaluation.yaml", "ch05_evaluation", "Chapter-05-Evaluation.docx"),
    "6": ("ch06_conclusion.yaml", "ch06_conclusion", "Chapter-06-Conclusion.docx"),
}


def _meta_from_cfg(meta_cfg: dict) -> dict:
    return {
        "font": meta_cfg.get("font_primary", "Iskoola Pota"),
        "font_latin": meta_cfg.get("font_latin", "Times New Roman"),
        "font_size": meta_cfg.get("font_size_body", 12),
        "line_spacing": meta_cfg.get("line_spacing", 1.5),
    }


def _blank_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(4.0)
    sec.right_margin = Cm(2.54)
    return doc


def build_chapter(chapter_num: str) -> tuple[Document, dict, Path]:
    yaml_name, bulk_key, filename = CHAPTER_MAP[chapter_num]
    meta_cfg = gt.load_yaml("meta.yaml")
    bulk = gt.load_yaml("bulk.yaml")
    figures_cfg = gt.load_yaml("figures.yaml")
    chapter = gt.load_yaml(yaml_name)
    gt.merge_bulk(chapter, bulk, bulk_key)
    meta = _meta_from_cfg(meta_cfg)

    doc = _blank_doc()
    gt.add_heading(doc, chapter.get("title", ""), 1, meta)
    for section in chapter.get("sections") or []:
        gt.render_section(doc, section, meta, figures_cfg.get("figures") or {})

    return doc, meta, ROOT / "docs" / "thesis" / filename


def build_appendices() -> tuple[Document, dict, Path]:
    """One Word file: Appendices A–E, then References (post–Chapter 6 material)."""
    meta_cfg = gt.load_yaml("meta.yaml")
    figures_cfg = gt.load_yaml("figures.yaml")
    supp = gt.load_yaml("supplementary.yaml")
    refs = gt.load_yaml("references.yaml")
    meta = _meta_from_cfg(meta_cfg)

    doc = _blank_doc()
    gt.add_heading(doc, "පරිශිෂ්ට (Appendices)", 1, meta)
    for section in supp.get("sections") or []:
        gt.render_section(doc, section, meta, figures_cfg.get("figures") or {})

    doc.add_page_break()
    gt.add_heading(doc, "මූලාශ්‍ර (References)", 1, meta)
    for ref in refs.get("items") or []:
        gt.add_reference(doc, ref, meta)

    return doc, meta, ROOT / "docs" / "thesis" / "Appendices-and-References.docx"


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate a standalone thesis chapter Word file.")
    parser.add_argument(
        "chapter",
        nargs="?",
        default="1",
        help="Chapter number (1-6), or 'appendices' / 'A' for Appendices+References",
    )
    args = parser.parse_args()
    key = args.chapter.strip().lower()

    if key in {"appendices", "appendix", "a", "7", "refs"}:
        doc, meta, out = build_appendices()
    elif args.chapter in CHAPTER_MAP:
        doc, meta, out = build_chapter(args.chapter)
    else:
        print(
            f"Unknown chapter: {args.chapter}. "
            f"Choose from: {', '.join(CHAPTER_MAP)} or 'appendices'"
        )
        return 1

    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out))
    except PermissionError:
        out = out.with_name(out.stem + "-updated.docx")
        doc.save(str(out))
        print(f"NOTE: original file open in Word - wrote {out.name} instead.")

    words = gt.count_words(doc)
    print(f"Wrote {out}")
    print(f"Font: {meta['font']} {meta['font_size']}pt, line spacing {meta['line_spacing']}")
    print(f"Words: {words:,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
